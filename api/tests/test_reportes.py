"""Reporte del contador: totales de servicios por persona (issue #24).
 
No pasa por el flujo completo de crear/cerrar una atención con su
formato (eso ya está cubierto en `test_atenciones.py` y
`test_formato_*`) -- acá solo importa la fila de `Atencion` (quién fue
conductor, quién responsable, cuándo se cerró), así que se crea
directo en la base de datos de prueba.
"""
 
from __future__ import annotations
 
from datetime import UTC, date, datetime, time
 
import openpyxl
from app.models.ambulancia import Ambulancia, TipoAmbulancia
from app.models.atencion import Atencion, EstadoAtencion, TipoAtencion
from app.models.empleado import Empleado
from app.models.formato_accidente_soat import FormatoAccidenteSoat, RolVictima
from app.models.formato_traslado import (
    CategoriaPaciente,
    ComplejidadTraslado,
    FormatoTraslado,
    ModalidadTraslado,
    NivelServicioTraslado,
    SexoPaciente,
    TipoDocumentoPaciente,
)
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
 
 
def _crear_ambulancia(db: Session, movil: str = "M-rep") -> Ambulancia:
    ambulancia = Ambulancia(
        movil=movil,
        placa=f"{movil}PLACA"[:10],
        tipo=TipoAmbulancia.basica,
        vencimiento_soat=date(2027, 1, 1),
        vencimiento_tecnomecanica=date(2027, 1, 1),
        activa=True,
    )
    db.add(ambulancia)
    db.commit()
    db.refresh(ambulancia)
    return ambulancia
 
 
def _crear_empleado(db: Session, nombres: str, apellidos: str, cedula: str) -> Empleado:
    empleado = Empleado(nombres=nombres, apellidos=apellidos, cedula=cedula)
    db.add(empleado)
    db.commit()
    db.refresh(empleado)
    return empleado
 
 
def _crear_atencion(
    db: Session,
    ambulancia: Ambulancia,
    conductor: Empleado,
    responsable: Empleado,
    *,
    estado: EstadoAtencion = EstadoAtencion.cerrado,
    cerrada_en: datetime | None = None,
    tipo: TipoAtencion = TipoAtencion.traslado,
    nivel_servicio: NivelServicioTraslado = NivelServicioTraslado.basico,
    paciente_nombres: str = "Paciente",
    paciente_apellidos: str = "De Prueba",
    paciente_numero_documento: str = "1000000000",
    recepcion_ips: str = "IPS de recepción",
    recepcion_ciudad: str = "Ocaña",
    entrega_ips: str = "IPS de entrega",
    entrega_ciudad: str = "Cúcuta",
    lugar_ocurrencia: str = "Calle 10 con carrera 5",
    ciudad_soat: str = "Ocaña",
    identificado_soat: bool = True,
) -> Atencion:
    """Crea la atención y, según el tipo, también el encabezado del
    formato correspondiente (`FormatoTraslado` o `FormatoAccidenteSoat`)
    -- en el sistema real siempre existe para una atención cerrada (se
    llena al recibir al paciente, no al cerrar), así que tanto el
    desglose por tipo (issue #24) como el drill-down de viajes
    (issue #25) lo asumen ahí.
    """
    abierta_en = cerrada_en or datetime(2027, 1, 1, tzinfo=UTC)
    atencion = Atencion(
        ambulancia_id=ambulancia.id,
        conductor_id=conductor.id,
        responsable_id=responsable.id,
        tipo=tipo,
        estado=estado,
        abierta_en=abierta_en,
        cerrada_en=cerrada_en,
    )
    db.add(atencion)
    db.commit()
    db.refresh(atencion)
 
    if tipo == TipoAtencion.traslado:
        formato_traslado = FormatoTraslado(
            atencion_id=atencion.id,
            paciente_tipo_documento=TipoDocumentoPaciente.cc,
            paciente_numero_documento=paciente_numero_documento,
            paciente_eps="EPS de prueba",
            paciente_nombres=paciente_nombres,
            paciente_apellidos=paciente_apellidos,
            paciente_edad=30,
            paciente_sexo=SexoPaciente.m,
            paciente_direccion_residencial="Calle falsa 123",
            paciente_ciudad="Ocaña",
            recepcion_fecha=abierta_en.date(),
            recepcion_hora=time(8, 0),
            recepcion_ciudad=recepcion_ciudad,
            recepcion_ips=recepcion_ips,
            recepcion_servicio="Urgencias",
            entrega_fecha=abierta_en.date(),
            entrega_hora=time(9, 0),
            entrega_ciudad=entrega_ciudad,
            entrega_ips=entrega_ips,
            entrega_servicio="Urgencias",
            complejidad=ComplejidadTraslado.baja,
            categoria_paciente=CategoriaPaciente.adulto,
            nivel_servicio=nivel_servicio,
            modalidad=ModalidadTraslado.sencillo,
        )
        db.add(formato_traslado)
        db.commit()
    elif tipo == TipoAtencion.atencion_soat:
        formato_soat = FormatoAccidenteSoat(
            atencion_id=atencion.id,
            fecha_atencion=abierta_en.date(),
            hora_atencion=time(8, 0),
            identificado=identificado_soat,
            paciente_tipo_documento=TipoDocumentoPaciente.cc if identificado_soat else None,
            paciente_numero_documento=paciente_numero_documento if identificado_soat else None,
            paciente_nombres=paciente_nombres if identificado_soat else None,
            paciente_apellidos=paciente_apellidos if identificado_soat else None,
            paciente_edad=30 if identificado_soat else None,
            paciente_sexo=SexoPaciente.m if identificado_soat else None,
            lugar_ocurrencia=lugar_ocurrencia,
            ciudad=ciudad_soat,
            rol_victima=RolVictima.peaton,
        )
        db.add(formato_soat)
        db.commit()
 
    return atencion
 
 
class TestPermisos:
    def test_admin_puede_consultar(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
 
    def test_contador_puede_consultar(
        self, client: TestClient, db: Session, contador_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=contador_headers,
        )
        assert respuesta.status_code == 200
 
    def test_auxiliar_no_puede_consultar(
        self, client: TestClient, db: Session, auxiliar_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=auxiliar_headers,
        )
        assert respuesta.status_code == 403
 
    def test_conductor_no_puede_consultar(
        self, client: TestClient, db: Session, conductor_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=conductor_headers,
        )
        assert respuesta.status_code == 403
 
    def test_sin_sesion_da_401(self, client: TestClient) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
        )
        assert respuesta.status_code == 401
 
 
class TestValidacionDeRango:
    def test_desde_mayor_que_hasta_da_422(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-31", "hasta": "2027-01-01"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 422
 
    def test_faltan_parametros_da_422(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get("/reportes/servicios-por-persona", headers=admin_headers)
        assert respuesta.status_code == 422
 
 
class TestTotales:
    def test_vacio_si_no_hay_cerradas_en_el_rango(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
        assert respuesta.json() == []
 
    def test_cuenta_conductor_y_responsable_en_filas_separadas(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000100")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000100")
        # Dos servicios: mismo conductor y mismo auxiliar en ambos.
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 1, 5, tzinfo=UTC)
        )
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 1, 10, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
        filas = respuesta.json()
        assert len(filas) == 2
 
        por_cedula = {fila["cedula"]: fila for fila in filas}
        assert por_cedula["800000100"]["funcion"] == "conductor"
        assert por_cedula["800000100"]["total"] == 2
        assert por_cedula["700000100"]["funcion"] == "responsable"
        assert por_cedula["700000100"]["total"] == 2
 
    def test_una_persona_que_hace_de_ambas_cosas_sale_en_dos_filas(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        """Decisión del cliente (mockup compartido en issue #24): no
        se combina en una sola fila -- una fila por persona y función.
        """
        ambulancia = _crear_ambulancia(db)
        empleado = _crear_empleado(db, "Carlos", "Gómez", "800000101")
        otro = _crear_empleado(db, "Laura", "Gómez", "700000101")
        # Un servicio donde `empleado` conduce, otro donde es responsable.
        _crear_atencion(
            db, ambulancia, empleado, otro, cerrada_en=datetime(2027, 1, 5, tzinfo=UTC)
        )
        _crear_atencion(
            db, ambulancia, otro, empleado, cerrada_en=datetime(2027, 1, 6, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        filas_empleado = [f for f in respuesta.json() if f["cedula"] == "800000101"]
        assert len(filas_empleado) == 2
        por_funcion = {f["funcion"]: f for f in filas_empleado}
        assert por_funcion["conductor"]["total"] == 1
        assert por_funcion["responsable"]["total"] == 1
 
    def test_desglosa_basicos_medicalizados_y_soat(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000107")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000107")
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 5, tzinfo=UTC),
            tipo=TipoAtencion.traslado,
            nivel_servicio=NivelServicioTraslado.basico,
        )
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 6, tzinfo=UTC),
            tipo=TipoAtencion.traslado,
            nivel_servicio=NivelServicioTraslado.medicalizado,
        )
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 7, tzinfo=UTC),
            tipo=TipoAtencion.atencion_soat,
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        fila = next(
            f
            for f in respuesta.json()
            if f["cedula"] == "800000107" and f["funcion"] == "conductor"
        )
        assert fila["basicos"] == 1
        assert fila["medicalizados"] == 1
        assert fila["soat"] == 1
        assert fila["total"] == 3
 
    def test_no_cuenta_atenciones_abiertas(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000102")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000102")
        _crear_atencion(db, ambulancia, conductor, auxiliar, estado=EstadoAtencion.abierto)
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.json() == []
 
    def test_no_cuenta_cerradas_fuera_del_rango(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000103")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000103")
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 2, 1, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.json() == []
 
    def test_limites_del_rango_son_inclusivos(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000104")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000104")
        # Justo al principio y al final del día del límite.
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 1, 0, 0, 1, tzinfo=UTC),
        )
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 31, 23, 59, 59, tzinfo=UTC),
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        fila = next(f for f in respuesta.json() if f["cedula"] == "800000104")
        assert fila["total"] == 2
 
    def test_todos_los_tipos_de_atencion_cuentan_igual(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        """Decisión del cliente: traslado y accidente SOAT cuentan
        igual como servicio para el pago (ver nota de módulo de
        `app/api/routes/reportes.py` y la decisión guardada del
        issue #24)."""
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000105")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000105")
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 5, tzinfo=UTC),
            tipo=TipoAtencion.traslado,
        )
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 6, tzinfo=UTC),
            tipo=TipoAtencion.atencion_soat,
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        fila = next(f for f in respuesta.json() if f["cedula"] == "800000105")
        assert fila["total"] == 2
 
 
class TestExcel:
    def test_contador_puede_descargar(
        self, client: TestClient, db: Session, contador_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona.xlsx",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=contador_headers,
        )
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "servicios-por-persona_2027-01-01_2027-01-31.xlsx" in (
            respuesta.headers["content-disposition"]
        )
 
    def test_auxiliar_no_puede_descargar(
        self, client: TestClient, db: Session, auxiliar_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona.xlsx",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=auxiliar_headers,
        )
        assert respuesta.status_code == 403
 
    def test_desde_mayor_que_hasta_da_422(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona.xlsx",
            params={"desde": "2027-01-31", "hasta": "2027-01-01"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 422
 
    def test_contenido_del_excel_coincide_con_el_json(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000106")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000106")
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 1, 5, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona.xlsx",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
 
        from io import BytesIO
 
        libro = openpyxl.load_workbook(BytesIO(respuesta.content))
        hoja = libro.active
        encabezados = [celda.value for celda in next(hoja.iter_rows(min_row=1, max_row=1))]
        assert encabezados == [
            "Cédula",
            "Nombres",
            "Apellidos",
            "Función",
            "Básicos",
            "Medicalizados",
            "SOAT",
            "Total",
        ]
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        # Un servicio básico: una fila para el conductor, otra para el
        # auxiliar, más la fila de totales al final.
        assert ("800000106", "Pedro", "Londoño", "Conductor", 1, 0, 0, 1) in filas
        assert ("700000106", "Ana", "Ruiz", "Auxiliar/médico", 1, 0, 0, 1) in filas
        assert filas[-1] == (None, None, None, "Total", 2, 0, 0, 2)
 
    def test_excel_sin_filas_no_trae_fila_de_totales(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona.xlsx",
            params={"desde": "2027-01-01", "hasta": "2027-01-31"},
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
 
        from io import BytesIO
 
        libro = openpyxl.load_workbook(BytesIO(respuesta.content))
        hoja = libro.active
        assert list(hoja.iter_rows(min_row=2, values_only=True)) == []
 
 
class TestDetalle:
    """Drill-down de viajes por persona y función (issue #25)."""
 
    def test_admin_puede_consultar(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": 1,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
 
    def test_contador_puede_consultar(
        self, client: TestClient, db: Session, contador_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": 1,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=contador_headers,
        )
        assert respuesta.status_code == 200
 
    def test_auxiliar_no_puede_consultar(
        self, client: TestClient, db: Session, auxiliar_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": 1,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=auxiliar_headers,
        )
        assert respuesta.status_code == 403
 
    def test_funcion_invalida_da_422(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": 1,
                "funcion": "mecanico",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        assert respuesta.status_code == 422
 
    def test_trae_origen_destino_paciente_cedula_y_fecha_de_un_traslado(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000200")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000200")
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 5, tzinfo=UTC),
            paciente_nombres="Juan",
            paciente_apellidos="Pérez",
            paciente_numero_documento="900000001",
            recepcion_ips="Clínica Norte",
            recepcion_ciudad="Ocaña",
            entrega_ips="Hospital Central",
            entrega_ciudad="Cúcuta",
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": conductor.id,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        assert respuesta.status_code == 200
        viajes = respuesta.json()
        assert len(viajes) == 1
        assert viajes[0]["fecha"] == "2027-01-05"
        assert viajes[0]["origen"] == "Clínica Norte (Ocaña)"
        assert viajes[0]["destino"] == "Hospital Central (Cúcuta)"
        assert viajes[0]["paciente"] == "Juan Pérez"
        assert viajes[0]["cedula"] == "900000001"
 
    def test_soat_no_tiene_destino_y_puede_ser_nn(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000201")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000201")
        _crear_atencion(
            db,
            ambulancia,
            conductor,
            auxiliar,
            cerrada_en=datetime(2027, 1, 6, tzinfo=UTC),
            tipo=TipoAtencion.atencion_soat,
            lugar_ocurrencia="Avenida Circunvalar",
            ciudad_soat="Ocaña",
            identificado_soat=False,
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": conductor.id,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        viajes = respuesta.json()
        assert len(viajes) == 1
        assert viajes[0]["origen"] == "Avenida Circunvalar (Ocaña)"
        assert viajes[0]["destino"] == "No aplica"
        assert viajes[0]["paciente"] == "NN"
        assert viajes[0]["cedula"] == ""
 
    def test_solo_trae_los_viajes_de_la_funcion_pedida(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        empleado = _crear_empleado(db, "Carlos", "Gómez", "800000202")
        otro = _crear_empleado(db, "Laura", "Gómez", "700000202")
        # `empleado` conduce un servicio y es responsable de otro --
        # el drill-down de "conductor" solo debe traer el primero.
        _crear_atencion(
            db, ambulancia, empleado, otro, cerrada_en=datetime(2027, 1, 5, tzinfo=UTC)
        )
        _crear_atencion(
            db, ambulancia, otro, empleado, cerrada_en=datetime(2027, 1, 6, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": empleado.id,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        viajes = respuesta.json()
        assert len(viajes) == 1
        assert viajes[0]["fecha"] == "2027-01-05"
 
    def test_no_trae_viajes_fuera_del_rango(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000203")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000203")
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 2, 1, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": conductor.id,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        assert respuesta.json() == []
 
    def test_viajes_ordenados_por_fecha(
        self, client: TestClient, db: Session, admin_headers: dict[str, str]
    ) -> None:
        ambulancia = _crear_ambulancia(db)
        conductor = _crear_empleado(db, "Pedro", "Londoño", "800000204")
        auxiliar = _crear_empleado(db, "Ana", "Ruiz", "700000204")
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 1, 20, tzinfo=UTC)
        )
        _crear_atencion(
            db, ambulancia, conductor, auxiliar, cerrada_en=datetime(2027, 1, 5, tzinfo=UTC)
        )
 
        respuesta = client.get(
            "/reportes/servicios-por-persona/detalle",
            params={
                "empleado_id": conductor.id,
                "funcion": "conductor",
                "desde": "2027-01-01",
                "hasta": "2027-01-31",
            },
            headers=admin_headers,
        )
        fechas = [v["fecha"] for v in respuesta.json()]
        assert fechas == ["2027-01-05", "2027-01-20"]