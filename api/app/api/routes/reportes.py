"""Reporte del contador: totales de servicios por persona en un rango
de fechas (issue #24) -- para el tema del pago, cuántos servicios
básicos, medicalizados y SOAT hizo cada persona, separado por si
participó como conductor o como auxiliar/médico responsable, con
descarga en Excel.
 
Solo cuentan las atenciones ya CERRADAS: una recién abierta todavía
puede cambiar o cancelarse, así que no es un servicio "hecho" para
efectos de pago (decisión del cliente, issue #24). El rango filtra por
`cerrada_en` -- la fecha en la que el servicio realmente se completó,
mismo campo que ya usa la lista de servicios cerrados del panel
(`GET /atenciones/cerradas`, issue #15/#80) -- no `abierta_en`, que
podría quedar fuera del rango si el cierre pasó un tiempo después.
 
`desde`/`hasta` son fechas de calendario (sin hora): igual que el
resto del proyecto, que guarda todo en UTC sin manejar explícitamente
la zona horaria de Colombia (ver `datetime.now(UTC)` en todo
`atenciones.py`), el rango se interpreta como ese día calendario en
UTC -- no hay conversión especial acá tampoco.
 
Columnas del desglose (mockup compartido por el cliente): "básico" y
"medicalizado" son los dos niveles del formato de traslado
(`FormatoTraslado.nivel_servicio`, siempre presente en una atención de
tipo traslado -- se llena al recibir al paciente, no al cerrar); SOAT
es el otro tipo de atención (`TipoAtencion.atencion_soat`), que no
tiene esa subdivisión y cuenta entero en su propia columna.
 
Cada persona aparece en una fila por función (conductor/responsable):
si alguien fue conductor en unos servicios y auxiliar/médico
responsable en otros dentro del rango, son dos filas distintas, no una
combinada -- también decisión explícita del cliente (issue #24).
 
Issue #25 agrega el drill-down: al hacer clic en una fila, el detalle
de cada viaje de esa persona en esa función (origen, destino,
paciente, cédula, fecha) -- sin el formato completo ni datos clínicos,
a pedido explícito del cliente. El formato de traslado tiene origen y
destino bien definidos (recepción/entrega); el de accidente SOAT solo
tiene el lugar del accidente (no existe "destino" en ese papel), así
que esa columna queda como "No aplica" para esos viajes. El paciente
de un accidente SOAT puede ser NN (sin identificar) -- se muestra así,
con la cédula vacía.
"""
 
from __future__ import annotations
 
from collections import defaultdict
from datetime import UTC, date, datetime, time
from io import BytesIO
 
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload
 
from app.api.deps import DbSession, require_admin_o_contador
from app.models.atencion import Atencion, EstadoAtencion, TipoAtencion
from app.models.empleado import Empleado
from app.models.formato_accidente_soat import FormatoAccidenteSoat
from app.models.formato_traslado import FormatoTraslado, NivelServicioTraslado
from app.schemas.reporte import Funcion, ServicioPorPersona, ViajeDetalle
 
router = APIRouter(
    prefix="/reportes",
    tags=["reportes"],
    dependencies=[Depends(require_admin_o_contador)],
)
 
_ENCABEZADOS = (
    "Cédula",
    "Nombres",
    "Apellidos",
    "Función",
    "Básicos",
    "Medicalizados",
    "SOAT",
    "Total",
)
_ANCHOS_COLUMNA = (16, 20, 20, 14, 10, 14, 8, 10)
_ETIQUETA_FUNCION: dict[Funcion, str] = {"conductor": "Conductor", "responsable": "Auxiliar/médico"}
 
 
def _validar_rango(desde: date, hasta: date) -> None:
    if desde > hasta:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="'desde' no puede ser posterior a 'hasta'",
        )
 
 
def _categoria(atencion: Atencion, niveles: dict[int, NivelServicioTraslado]) -> str:
    """A cuál de las tres columnas del reporte (básicos/medicalizados/
    soat) suma esta atención.
    """
    if atencion.tipo == TipoAtencion.atencion_soat:
        return "soat"
    nivel = niveles[atencion.id]
    return "medicalizados" if nivel == NivelServicioTraslado.medicalizado else "basicos"
 
 
def _calcular_totales(db: Session, desde: date, hasta: date) -> list[ServicioPorPersona]:
    """Trae las atenciones cerradas cuyo `cerrada_en` cae en
    `[desde, hasta]` (límites inclusive, día calendario completo) y
    arma los totales por (empleado, función) en Python -- no son
    tantas filas por rango típico (un cierre de nómina quincenal/
    mensual de 4 ambulancias) como para justificar el `GROUP BY` en
    SQL.
    """
    inicio = datetime.combine(desde, time.min, tzinfo=UTC)
    fin = datetime.combine(hasta, time.max, tzinfo=UTC)
    atenciones = (
        db.execute(
            select(Atencion)
            .options(joinedload(Atencion.conductor), joinedload(Atencion.responsable))
            .where(
                Atencion.estado == EstadoAtencion.cerrado,
                Atencion.cerrada_en >= inicio,
                Atencion.cerrada_en <= fin,
            )
        )
        .unique()
        .scalars()
        .all()
    )
 
    # El nivel de servicio (básico/medicalizado) vive en el formato de
    # traslado, no en la atención misma -- se trae en un solo query
    # para las atenciones de tipo traslado del rango, en vez de uno
    # por atención.
    ids_traslado = [a.id for a in atenciones if a.tipo == TipoAtencion.traslado]
    niveles: dict[int, NivelServicioTraslado] = {
        fila.atencion_id: fila.nivel_servicio
        for fila in db.execute(
            select(FormatoTraslado.atencion_id, FormatoTraslado.nivel_servicio).where(
                FormatoTraslado.atencion_id.in_(ids_traslado)
            )
        )
    }
 
    empleados: dict[tuple[int, Funcion], Empleado] = {}
    conteos: dict[tuple[int, Funcion], dict[str, int]] = defaultdict(
        lambda: {"basicos": 0, "medicalizados": 0, "soat": 0}
    )
    for atencion in atenciones:
        categoria = _categoria(atencion, niveles)
        clave_conductor: tuple[int, Funcion] = (atencion.conductor_id, "conductor")
        empleados[clave_conductor] = atencion.conductor
        conteos[clave_conductor][categoria] += 1
 
        clave_responsable: tuple[int, Funcion] = (atencion.responsable_id, "responsable")
        empleados[clave_responsable] = atencion.responsable
        conteos[clave_responsable][categoria] += 1
 
    filas = [
        ServicioPorPersona(
            empleado_id=empleado_id,
            cedula=empleado.cedula,
            nombres=empleado.nombres,
            apellidos=empleado.apellidos,
            funcion=funcion,
            basicos=conteos[(empleado_id, funcion)]["basicos"],
            medicalizados=conteos[(empleado_id, funcion)]["medicalizados"],
            soat=conteos[(empleado_id, funcion)]["soat"],
            total=sum(conteos[(empleado_id, funcion)].values()),
        )
        for (empleado_id, funcion), empleado in empleados.items()
    ]
    filas.sort(key=lambda fila: (fila.apellidos.lower(), fila.nombres.lower(), fila.funcion))
    return filas
 
 
def _armar_libro(filas: list[ServicioPorPersona]) -> Worksheet:
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None  # `Workbook()` siempre trae una hoja activa.
    hoja.title = "Servicios por persona"
    hoja.append(_ENCABEZADOS)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        hoja.append(
            (
                fila.cedula,
                fila.nombres,
                fila.apellidos,
                _ETIQUETA_FUNCION[fila.funcion],
                fila.basicos,
                fila.medicalizados,
                fila.soat,
                fila.total,
            )
        )
    if filas:
        hoja.append(
            (
                "",
                "",
                "",
                "Total",
                sum(f.basicos for f in filas),
                sum(f.medicalizados for f in filas),
                sum(f.soat for f in filas),
                sum(f.total for f in filas),
            )
        )
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
    for columna, ancho in zip(hoja.columns, _ANCHOS_COLUMNA, strict=True):
        hoja.column_dimensions[columna[0].column_letter].width = ancho
    return hoja
 
 
@router.get("/servicios-por-persona", response_model=list[ServicioPorPersona])
def reporte_servicios_por_persona(
    db: DbSession,
    desde: date,
    hasta: date,
) -> list[ServicioPorPersona]:
    """Totales de servicios cerrados por persona y función (conductor/
    responsable), desglosados por tipo, en `[desde, hasta]`, para la
    pantalla del reporte en el panel (issue #24).
    """
    _validar_rango(desde, hasta)
    return _calcular_totales(db, desde, hasta)
 
 
@router.get("/servicios-por-persona.xlsx")
def reporte_servicios_por_persona_excel(
    db: DbSession,
    desde: date,
    hasta: date,
) -> StreamingResponse:
    """Igual que `reporte_servicios_por_persona`, como archivo de
    Excel descargable -- el pedido explícito del cliente para el
    issue #24, no solo verlo en pantalla. Incluye una fila de totales
    al final (mockup del cliente).
    """
    _validar_rango(desde, hasta)
    filas = _calcular_totales(db, desde, hasta)
    hoja = _armar_libro(filas)
 
    buffer = BytesIO()
    hoja.parent.save(buffer)
    buffer.seek(0)
 
    nombre_archivo = f"servicios-por-persona_{desde.isoformat()}_{hasta.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
 
 
def _detalle_traslado(atencion: Atencion, formato: FormatoTraslado) -> ViajeDetalle:
    return ViajeDetalle(
        atencion_id=atencion.id,
        fecha=atencion.cerrada_en.date() if atencion.cerrada_en else atencion.abierta_en.date(),
        origen=f"{formato.recepcion_ips} ({formato.recepcion_ciudad})",
        # `entrega_ips`/`entrega_ciudad` son obligatorios para cerrar un
        # traslado (ver `cerrar_atencion` en `atenciones.py`), así que
        # siempre están presentes acá -- no hace falta un "No aplica".
        destino=f"{formato.entrega_ips} ({formato.entrega_ciudad})",
        paciente=f"{formato.paciente_nombres} {formato.paciente_apellidos}",
        cedula=formato.paciente_numero_documento,
    )
 
 
def _detalle_soat(atencion: Atencion, formato: FormatoAccidenteSoat) -> ViajeDetalle:
    return ViajeDetalle(
        atencion_id=atencion.id,
        fecha=atencion.cerrada_en.date() if atencion.cerrada_en else atencion.abierta_en.date(),
        origen=f"{formato.lugar_ocurrencia} ({formato.ciudad})",
        # El formato AA-LAFS-005 no tiene casilla de destino (solo el
        # lugar del accidente) -- decisión del cliente para el issue #25.
        destino="No aplica",
        # El paciente de un accidente puede llegar sin identificar
        # (NN) -- también decisión del cliente para este issue.
        paciente=(
            f"{formato.paciente_nombres or ''} {formato.paciente_apellidos or ''}".strip()
            if formato.identificado
            else "NN"
        ),
        cedula=(formato.paciente_numero_documento or "") if formato.identificado else "",
    )
 
 
def _calcular_detalle(
    db: Session, empleado_id: int, funcion: Funcion, desde: date, hasta: date
) -> list[ViajeDetalle]:
    """Los viajes de un empleado, en una función puntual, en el rango
    -- el drill-down del issue #25. Misma ventana de `[desde, hasta]`
    (sobre `cerrada_en`) que el resto del reporte.
    """
    inicio = datetime.combine(desde, time.min, tzinfo=UTC)
    fin = datetime.combine(hasta, time.max, tzinfo=UTC)
    columna_funcion = Atencion.conductor_id if funcion == "conductor" else Atencion.responsable_id
    atenciones = (
        db.execute(
            select(Atencion)
            .where(
                Atencion.estado == EstadoAtencion.cerrado,
                Atencion.cerrada_en >= inicio,
                Atencion.cerrada_en <= fin,
                columna_funcion == empleado_id,
            )
            .order_by(Atencion.cerrada_en)
        )
        .scalars()
        .all()
    )
 
    ids_traslado = [a.id for a in atenciones if a.tipo == TipoAtencion.traslado]
    ids_soat = [a.id for a in atenciones if a.tipo == TipoAtencion.atencion_soat]
    formatos_traslado = {
        f.atencion_id: f
        for f in db.execute(
            select(FormatoTraslado).where(FormatoTraslado.atencion_id.in_(ids_traslado))
        ).scalars()
    }
    formatos_soat = {
        f.atencion_id: f
        for f in db.execute(
            select(FormatoAccidenteSoat).where(FormatoAccidenteSoat.atencion_id.in_(ids_soat))
        ).scalars()
    }
 
    viajes = []
    for atencion in atenciones:
        if atencion.tipo == TipoAtencion.traslado:
            viajes.append(_detalle_traslado(atencion, formatos_traslado[atencion.id]))
        else:
            viajes.append(_detalle_soat(atencion, formatos_soat[atencion.id]))
    return viajes
 
 
@router.get("/servicios-por-persona/detalle", response_model=list[ViajeDetalle])
def reporte_detalle_servicios_por_persona(
    db: DbSession,
    empleado_id: int,
    funcion: Funcion,
    desde: date,
    hasta: date,
) -> list[ViajeDetalle]:
    """El drill-down de una fila del reporte (issue #25): cada viaje
    de ese empleado en esa función, con origen, destino, paciente,
    cédula y fecha -- sin el formato completo ni datos clínicos.
    """
    _validar_rango(desde, hasta)
    return _calcular_detalle(db, empleado_id, funcion, desde, hasta)